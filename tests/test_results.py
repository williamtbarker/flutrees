import json
from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock
import pandas as pd
import pytest
from pypdf import PdfReader
from openpyxl import load_workbook
from bs4 import BeautifulSoup

from flutrees import pipeline
from flutrees.exports import flat_mutations, write_pdf, write_workbook, write_html
from flutrees.figures import tree_pages, tree_figure, write_tree_figures
from flutrees.tree_build import build_tree, iter_nodes


def test_complete_run_and_artifact_contents(tmp_path, cfg, fasta):
    progress = []
    root = pipeline.run_many(cfg, [fasta], tmp_path / "results", "test", progress.append)
    out = root / "sample"
    status = json.loads((out / "status.json").read_text())
    assert status["status"] == "complete"
    assert json.loads((root / "status.json").read_text())["status"] == "complete"
    assert "Open this file:" in progress[-1]
    summary = json.loads((out / "summary.json").read_text())
    assert summary["n_records"] == 4 and len(summary["input_sha256"]) == 64
    assert summary["top_mutations"][0]["mutations"] == "F5Y"
    pdf = PdfReader(out / "report.pdf")
    assert len(pdf.pages) == 3
    assert "Pruned mutation" in pdf.pages[2].extract_text()
    assert "2 records" in pdf.pages[2].extract_text()
    wb = load_workbook(out / "results.xlsx")
    assert wb.sheetnames == [
        "Summary",
        "Records",
        "Mutations",
        "Nodes",
        "Node Membership",
        "QC",
        "Parameters",
    ]
    assert wb["Records"].max_row == 5 and wb["Records"].freeze_panes == "A2"
    assert len(wb["Records"].tables) == 1
    members = pd.read_csv(out / "node_membership.tsv", sep="\t")
    assert set(members.query("view=='full' and node_id==1").record_id) == set("abcd")
    soup = BeautifulSoup((out / "START_HERE.html").read_text(), "html.parser")
    assert soup.find("script") is None and soup.find("details")
    for a in soup.find_all("a", href=True):
        if not a["href"].startswith("#"):
            assert (out / a["href"]).exists()
    for name in ["tree_full.svg", "tree_pruned.svg"]:
        svg = BeautifulSoup((out / name).read_text(), "xml")
        assert len(svg.find_all("text")) >= 6
    assert (out / "tree_full.png").read_bytes().startswith(b"\x89PNG")
    before = (out / "summary.json").read_bytes()
    with pytest.raises(FileExistsError):
        pipeline.run_many(cfg, [fasta], tmp_path / "results", "test")
    assert (out / "summary.json").read_bytes() == before


def test_preflight_errors(tmp_path, cfg, fasta):
    for name in ["", ".", "..", "../escape", "a/b", "a\\b", "C:drive"]:
        with pytest.raises(ValueError, match="Run ID"):
            pipeline.run_many(cfg, [fasta], tmp_path, name)
    with pytest.raises(ValueError, match="at least"):
        pipeline.run_many(cfg, [], tmp_path, "test")
    with pytest.raises(ValueError, match="distinct name"):
        pipeline.run_many(cfg, [fasta, fasta], tmp_path, "test")
    with pytest.raises(ValueError):
        pipeline.input_name(Path(".."))
    invalid = tmp_path / "bad.fa"
    invalid.write_text("")
    with pytest.raises(ValueError):
        pipeline.run_many(cfg, [fasta, invalid], tmp_path / "preflight", "test")
    assert not (tmp_path / "preflight").exists()


def test_failed_exports_and_empty_outputs_are_not_success(tmp_path, cfg, fasta, monkeypatch):
    monkeypatch.setattr("flutrees.pipeline.write_pdf", Mock(side_effect=OSError("disk problem")))
    with pytest.raises(OSError):
        pipeline.run_many(cfg, [fasta], tmp_path / "results", "broken")
    root = tmp_path / "results" / "broken"
    assert json.loads((root / "status.json").read_text())["status"] == "failed"
    assert json.loads((root / "sample" / "status.json").read_text())["status"] == "failed"
    assert not (root / "START_HERE.html").exists()
    monkeypatch.setattr("flutrees.pipeline._analyze", lambda *a: None)
    with pytest.raises(ValueError, match="missing or empty"):
        pipeline.run_one(cfg, fasta, tmp_path / "empty")
    assert (
        json.loads((tmp_path / "empty" / "sample" / "status.json").read_text())["status"]
        == "failed"
    )


def test_alignment_id_validation(tmp_path, cfg, fasta, monkeypatch):
    monkeypatch.setattr(
        "flutrees.pipeline.mafft_align", lambda a, b, *args: b.write_text(">wrong\nACDEFG\n")
    )
    with pytest.raises(ValueError, match="IDs"):
        pipeline.run_one(cfg, fasta, tmp_path / "wrong")


def test_qc_identical_missing_duplicates_and_pruning(tmp_path, cfg):
    fasta = tmp_path / "qc.fa"
    fasta.write_text(">a\nACDEFG\n>a\nACD\n>c\nACDEFG\n>d\nACDEYG\n")
    root = pipeline.run_one(replace(cfg, prune_cutoff=3), fasta, tmp_path / "results")
    warnings = " ".join(json.loads((root / "summary.json").read_text())["warnings"])
    assert "ambiguous" in warnings and "padded" in warnings and "Duplicate" in warnings
    assert "No eligible split" in warnings
    fasta2 = tmp_path / "identical.fa"
    fasta2.write_text(">a\nACDEFG\n>b\nACDEFG\n")
    root = pipeline.run_one(cfg, fasta2, tmp_path / "results")
    assert "No substitutions" in " ".join(
        json.loads((root / "summary.json").read_text())["warnings"]
    )
    fasta3 = tmp_path / "prune.fa"
    fasta3.write_text(">a\nACDEFG\n>b\nACDEFG\n>c\nACDEYG\n>d\nACDEYG\n")
    root = pipeline.run_one(replace(cfg, prune_cutoff=3), fasta3, tmp_path / "results")
    assert "hides 2" in " ".join(json.loads((root / "summary.json").read_text())["warnings"])
    assert "hidden by pruning" in (root / "START_HERE.html").read_text()


def test_paginated_tree_and_figure_geometry(tmp_path):
    # Five independent substitutions generate a complete 5-level tree.
    data = pd.DataFrame(
        {
            "record_id": [str(i) for i in range(32)],
            "mutations": [
                [f"A{bit + 1}T" for bit in range(5) if i & (1 << bit)] for i in range(32)
            ],
        }
    )
    tree = build_tree(data, 6, 1, 0)
    terminal_page = tree_pages(build_tree(data, 3, 1, 0))
    assert len(terminal_page) == 1 and len(terminal_page[0]) == 15
    pages = tree_pages(tree)
    assert len(pages) == 9
    assert {n.node_id for p in pages for n in p} == {n.node_id for n in iter_nodes(tree)}
    mapping = {p[0].node_id: i for i, p in enumerate(pages, 1)}
    fig = tree_figure(pages[0], 32, "Test", 1, mapping)
    fig.canvas.draw()
    boxes = [t.get_window_extent(fig.canvas.get_renderer()) for t in fig.axes[0].texts]
    for i, b in enumerate(boxes):
        assert fig.bbox.contains(b.x0, b.y0) and fig.bbox.contains(b.x1, b.y1)
        assert all(not b.overlaps(other) for other in boxes[i + 1 :])
    import matplotlib.pyplot as plt

    plt.close(fig)
    write_tree_figures(tmp_path, tree, "tree_full")
    assert len(PdfReader(tmp_path / "tree_full.pdf").pages) == 9
    assert (tmp_path / "tree_full_page_009.svg").is_file()


def test_text_safety_and_long_summary(tmp_path):
    attack = "=SUM(1,2)<script>alert(1)</script>"
    df = pd.DataFrame({"record_id": [attack], "mutations": [[]]})
    tree = build_tree(df, 0, 1, 0)
    meta = pd.DataFrame({"record_id": [attack], "strain": [attack]})
    summary = {
        "input_name": attack,
        "n_records": 1,
        "reference_id": "r" * 2000,
        "start_residue": 1,
        "end_residue": 6,
        "warnings": [],
        "config": {"min_split": 1},
    }
    counts = pd.DataFrame({"mutations": pd.Series(dtype=str), "count": pd.Series(dtype=int)})
    write_html(tmp_path / "index.html", summary, tree, tree, meta)
    assert "<script>" not in (tmp_path / "index.html").read_text()
    write_workbook(tmp_path / "results.xlsx", summary, meta, df, counts, tree, tree)
    wb = load_workbook(tmp_path / "results.xlsx")
    assert wb["Records"]["A2"].data_type == "s" and wb["Records"]["A2"].value == attack
    write_pdf(tmp_path / "report.pdf", summary, counts, tree)
    assert len(PdfReader(tmp_path / "report.pdf").pages) >= 3
    assert flat_mutations(pd.DataFrame({"mutations": [None]})).iloc[0, 0] == ""
