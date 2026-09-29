"""Worksheet continuation and exact-value protection at format boundaries."""

import pandas as pd
import pytest
from openpyxl import load_workbook

from flutrees import exports
from flutrees.tree_build import build_tree


def test_worksheet_continuation_conserves_every_record(tmp_path, monkeypatch):
    monkeypatch.setattr(exports, "EXCEL_DATA_ROWS", 3)
    data = pd.DataFrame({"record_id": [str(i) for i in range(8)], "mutations": [[] for _ in range(8)]})
    tree = build_tree(data, 2, 1, 0)
    summary = {"input_name": "panel", "n_records": 8, "reference_id": "0", "start_residue": 1,
               "end_residue": 1, "config": {}, "warnings": []}
    extra = {"Table": pd.DataFrame({"id": range(8)})}
    exports.write_workbook(tmp_path / "out.xlsx", summary, data[["record_id"]], data,
                           pd.DataFrame({"mutation": []}), tree, tree, extra)
    book = load_workbook(tmp_path / "out.xlsx", data_only=True)
    assert [row[0].value for name in ("Records", "Records (2)", "Records (3)")
            for row in list(book[name].rows)[1:]] == data.record_id.tolist()
    assert [row[0].value for name in ("Table", "Table (2)", "Table (3)")
            for row in list(book[name].rows)[1:]] == list(range(8))
    assert all(sheet.freeze_panes == "A2" for sheet in book)
    assert all(sheet.max_row <= 4 for sheet in book if not sheet.title.startswith("Summary"))
    assert all(len(sheet.tables) == 1 for sheet in book if sheet.title.startswith("Records"))


def test_excel_never_silently_truncates_a_value(tmp_path):
    data = pd.DataFrame({"record_id": ["a"], "mutations": [[]]})
    tree = build_tree(data, 0, 1, 0)
    summary = {"input_name": "panel", "n_records": 1, "reference_id": "a", "start_residue": 1,
               "end_residue": 1, "config": {}, "warnings": []}
    metadata = pd.DataFrame({"record_id": ["a"], "raw_header": ["x" * 32_768]})
    with pytest.raises(ValueError, match="Excel cell limit"):
        exports.write_workbook(tmp_path / "long.xlsx", summary, metadata, data,
                               pd.DataFrame({"mutation": []}), tree, tree)


def test_real_excel_row_boundary(tmp_path):
    data = pd.DataFrame({"record_id": ["a"], "mutations": [[]]})
    tree = build_tree(data, 0, 1, 0)
    summary = {"input_name": "boundary", "n_records": 1, "reference_id": "a", "start_residue": 1,
               "end_residue": 1, "config": {}, "warnings": []}
    expected = exports.EXCEL_DATA_ROWS + 1
    output = tmp_path / "boundary.xlsx"
    exports.write_workbook(output, summary, data[["record_id"]], data,
                           pd.DataFrame({"mutation": []}), tree, tree,
                           {"Boundary": pd.DataFrame({"row_number": range(expected)})})
    book = load_workbook(output, read_only=True, data_only=True)
    try:
        assert book["Boundary"].max_row == 1_048_576
        assert book["Boundary (2)"].max_row == 2
        count = 0
        for name in ("Boundary", "Boundary (2)"):
            for row in book[name].iter_rows(min_row=2, values_only=True):
                assert row == (count,)
                count += 1
        assert count == expected
    finally:
        book.close()
