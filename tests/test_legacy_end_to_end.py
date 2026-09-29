"""Differential end-to-end acceptance against the pinned v0.2.3 source."""

import hashlib
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from flutrees.tree_io import read_tree
from flutrees.tree_build import to_dict


@pytest.fixture(scope="module")
def legacy_source():
    path = Path(os.environ.get("FLUTREES_LEGACY_SOURCE", "validation/legacy/src")).resolve()
    manifest = json.loads((Path(__file__).parent / "data/legacy_source_manifest.json").read_text())
    if not path.is_dir():
        pytest.fail("Check out v0.2.3 at validation/legacy, or set FLUTREES_LEGACY_SOURCE to its src directory. See CONTRIBUTING.md.")
    for relative, expected in manifest["files"].items():
        assert hashlib.sha256((path / relative).read_bytes()).hexdigest() == expected, relative
    return path


@pytest.mark.integration
@pytest.mark.skipif(shutil.which("mafft") is None, reason="Real MAFFT is required")
@pytest.mark.parametrize("fixture", ["synthetic", "public", "unknowns", "short-duplicates"])
@pytest.mark.parametrize("alignment", ["default", "legacy", "reproducible"])
def test_pinned_legacy_cli_scientific_outputs(tmp_path, legacy_source, fixture, alignment):
    repository = Path(__file__).resolve().parents[1]
    custom = []
    if fixture == "synthetic":
        source = repository / "src/flutrees/data/example.fasta"
    elif fixture == "public":
        source = repository / "tests/data/public_HA.fasta"
    else:
        source = tmp_path / "regression.fasta"
        if fixture == "unknowns":
            source.write_text(">duplicate\nAC?DEFGHIKLMNPQ*\n>duplicate\nAC?DEFGHIKLMNPQ*\n>variant\nAC?DEYGHIKLMNPQ*\n")
        else:
            source.write_text(">short\nACDEFG\n>short\nACDEFG\n>long\nACDEFGHIKLMNPQ\n>long\nACDEYGHIKLMNPQ\n")
        custom = ["--start", "1", "--end", "14", "--min-split", "1", "--prune-cutoff", "1"]
    for name, code in (("legacy", legacy_source), ("current", repository / "src")):
        env = os.environ.copy()
        env["PYTHONPATH"] = str(code)
        env.pop("COVERAGE_PROCESS_START", None)
        command = [sys.executable, "-m", "flutrees", "-i", str(source), "-o", str(tmp_path / name),
                   "--run-id", "comparison", "--threads", "1", *custom]
        if name == "current" and alignment != "default":
            command.append("--legacy-alignment" if alignment == "legacy" else "--reproducible")
        result = subprocess.run(command, env=env, cwd=tmp_path, capture_output=True, text=True, timeout=120)
        assert result.returncode == 0, result.stdout + result.stderr
    old = tmp_path / "legacy/comparison" / source.stem
    new = tmp_path / "current/comparison" / source.stem
    exact_artifacts = (
        "input.fasta", "extracted.fasta", "mafft_input.fasta", "aligned.fasta",
        "metadata.tsv", "mutations_per_record.tsv", "mutation_counts.tsv",
        "tree_full.json", "tree_pruned.json", "group_assignments.tsv", "node_summary.tsv", "node_membership.tsv",
    )
    for filename in exact_artifacts:
        assert (new / filename).read_bytes() == (old / filename).read_bytes(), filename
    for filename in ("tree_full.json", "tree_pruned.json"):
        assert to_dict(read_tree(old / filename)) == to_dict(read_tree(new / filename))
    before = json.loads((old / "summary.json").read_text())
    after = json.loads((new / "summary.json").read_text())
    for field in ("n_records", "start_residue", "end_residue", "reference_id", "reference_sequence",
                  "coordinate_system", "input_sha256", "top_mutations"):
        assert before[field] == after[field], field
    old_book = load_workbook(old / "results.xlsx", read_only=True)
    new_book = load_workbook(new / "results.xlsx", read_only=True)
    try:
        for sheet in ("Records", "Mutations", "Nodes", "Node Membership"):
            assert list(old_book[sheet].values) == list(new_book[sheet].values), sheet
    finally:
        old_book.close()
        new_book.close()
