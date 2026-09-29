"""End-to-end tree families, reference selection, portability, and provenance."""

import json
import subprocess
from dataclasses import replace
from unittest.mock import Mock

import pandas as pd
import pytest
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from pypdf import PdfReader
from typer.testing import CliRunner

from flutrees import pipeline
from flutrees.align_mafft import alignment_command
from flutrees.cli import app
from flutrees.config import RunConfig
from flutrees.exports import write_html
from flutrees.family import comparison_tables, expected_tree_files
from flutrees.provenance import mafft_version, portable_name, sha256_file
from flutrees.tree_build import build_tree


def test_all_modes_align_once_and_export_consistent_membership(tmp_path, cfg, fasta, monkeypatch):
    align = Mock(wraps=pipeline.mafft_align)
    monkeypatch.setattr(pipeline, "mafft_align", align)
    cfg = replace(cfg, tree_mode="all")
    out = pipeline.run_one(cfg, fasta, tmp_path / "results")
    align.assert_called_once()
    summary = json.loads((out / "summary.json").read_text())
    assert summary["tree_modes"] == ["frequency", "balanced", "diversity"]
    groups = pd.read_csv(out / "tree_groups.tsv", sep="\t")
    for mode in summary["tree_modes"]:
        mode_dir = out / "trees" / mode
        assert set(groups.loc[groups["mode"] == mode, "record_id"]) == set("abcd")
        assert json.loads((mode_dir / "method.json").read_text())["analysis_id"] == summary["analysis_id"]
        assert json.loads((mode_dir / "tree_full.json").read_text())["support"] == 4
    assert (out / "tree_full.json").read_bytes() == (out / "trees/frequency/tree_full.json").read_bytes()
    wb = load_workbook(out / "results.xlsx")
    assert "Tree Comparison" in wb and "Tree Groups" in wb and "All Membership" in wb
    assert wb["Tree Groups"].max_row == 13
    assert len(PdfReader(out / "report.pdf").pages) == 5
    html = BeautifulSoup((out / "START_HERE.html").read_text(), "html.parser")
    assert "Compare tree views" in html.get_text()
    for anchor in html.select("a[href]"):
        if not anchor["href"].startswith("#"):
            assert (out / anchor["href"]).is_file(), anchor["href"]
    manifest = json.loads((out / "status.json").read_text())
    assert manifest["artifacts"]["trees/diversity/tree_full.pdf"] > 0
    provenance = json.loads((out / "provenance.json").read_text())
    assert provenance["alignment_sha256"] == sha256_file(out / "aligned.fasta")
    assert provenance["mafft_command"][-3:-1] == ["--threadit", "0"]
    assert provenance["versions"]["python"] and len(provenance["analysis_id"]) == 64
    repeated = pipeline.run_one(cfg, fasta, tmp_path / "repeated")
    assert json.loads((repeated / "provenance.json").read_text())["analysis_id"] == provenance["analysis_id"]


def test_explicit_reference_and_alternative_cli(tmp_path, cfg, fasta):
    result = CliRunner().invoke(app, ["-i", str(fasta), "-o", str(tmp_path / "cli"), "--run-id", "ref",
                                    "--mafft", cfg.mafft, "--start", "1", "--end", "6", "--min-split", "1",
                                    "--tree-mode", "balanced", "--reference-id", "c", "--legacy-alignment"])
    assert result.exit_code == 0, result.output
    out = tmp_path / "cli/ref/sample"
    summary = json.loads((out / "summary.json").read_text())
    assert summary["reference_method"] == "explicit" and summary["reference_id"] == "c"
    assert summary["tree_strategy"] == "balanced"
    assert summary["top_mutations"][0]["mutations"] == "Y5F"
    assert "explicitly selected input record" in (out / "START_HERE.html").read_text()
    assert "--threadit" not in json.loads((out / "provenance.json").read_text())["mafft_command"]
    assert "Tree views: balanced" in result.output
    for ref in ["missing", "duplicate"]:
        file = tmp_path / f"{ref}.fasta"
        file.write_text(">duplicate\nACDEFG\n>duplicate\nACDEYG\n")
        with pytest.raises(ValueError, match="exactly one"):
            pipeline.run_many(replace(cfg, reference_id=ref), [file], tmp_path / ref, "fail")
        assert not (tmp_path / ref).exists()


def test_invalid_mode_reference_and_portable_run_id(tmp_path, cfg, fasta):
    for kwargs in [{"tree_mode": "invalid"}, {"reference_id": " "}]:
        with pytest.raises(ValueError):
            RunConfig(**kwargs)
    with pytest.raises(ValueError, match="portable"):
        pipeline.run_many(cfg, [fasta], tmp_path, "CON")
    for raw, expected in [("H1:panel", "H1_panel"), ("CON", "_CON"), ("NUL.txt", "_NUL.txt"),
                          ("trailing.", "trailing"), ("plain", "plain"), ("a?b", "a_b")]:
        assert portable_name(raw) == expected
    assert len(portable_name("漢" * 200).encode()) <= 120
    assert portable_name("a" * 200) != portable_name("a" * 199 + "b")
    for raw in ["", ".", ".."]:
        with pytest.raises(ValueError):
            portable_name(raw)
    first, second = tmp_path / "a:b.fasta", tmp_path / "a?b.fasta"
    for file in [first, second]:
        file.write_bytes(fasta.read_bytes())
    with pytest.raises(ValueError, match="distinct name"):
        pipeline.run_many(cfg, [first, second], tmp_path / "collision", "run")


def test_version_capture_errors_and_legacy_command(monkeypatch, tmp_path):
    runner = Mock(return_value=subprocess.CompletedProcess([], 0, stdout="v7.520\n"))
    monkeypatch.setattr("flutrees.provenance.subprocess.run", runner)
    assert mafft_version("mafft") == "v7.520"
    for stdout, code in [("", 0), ("failure", 1)]:
        runner.return_value = subprocess.CompletedProcess([], code, stdout=stdout)
        assert mafft_version("mafft") == "unavailable"
    for error in [OSError("gone"), subprocess.TimeoutExpired("mafft", 10)]:
        runner.side_effect = error
        assert mafft_version("mafft") == "unavailable"
    assert "--threadit" not in alignment_command("mafft", tmp_path / "a.fa", 1, False)


def test_family_failure_gate_and_root_only_tables(tmp_path, cfg, fasta, monkeypatch):
    original = pipeline._analyze

    def incomplete(*args):
        original(*args)
        (args[2] / "trees/frequency/method.json").unlink()

    monkeypatch.setattr(pipeline, "_analyze", incomplete)
    with pytest.raises(ValueError, match="tree-family"):
        pipeline.run_many(cfg, [fasta], tmp_path, "incomplete")
    assert json.loads((tmp_path / "incomplete/sample/status.json").read_text())["status"] == "failed"
    frame = pd.DataFrame({"record_id": ["a"], "mutations": [[]]})
    tree = build_tree(frame, 3, 1, 0)
    tables = comparison_tables({"frequency": (tree, tree)})
    assert tables["Mutation Use"].empty and tables["Tree Comparison"].iloc[0].root_split == "No split"
    summary = {"input_name": "one", "reference_id": "a", "n_records": 1,
               "start_residue": 1, "end_residue": 1, "warnings": []}
    write_html(tmp_path / "root.html", summary, tree, tree,
               pd.DataFrame({"record_id": ["a"], "strain": ["a"]}), family={"frequency": (tree, tree)})


def test_paginated_family_contract_covers_every_continuation():
    frame = pd.DataFrame({"record_id": list(map(str, range(32))),
                          "mutations": [[f"A{j+1}T" for j in range(5) if i & (1 << j)] for i in range(32)]})
    tree = build_tree(frame, 5, 1, 0)
    files = expected_tree_files(tree, "full")
    assert "tree_full_page_009.svg" in files and "tree_full_page_009.png" in files
    tables = comparison_tables({"frequency": (tree, tree)})
    assert tables["Mutation Use"].split_occurrences.max() > 1
